from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import ComparacaoItem, Documento, Indicador, Pendencia, ResultadoAuditoria


def _valor_enum(valor):
    return valor.value if valor is not None else None


def _pares_consecutivos(periodos: list[str]) -> list[tuple[str, str]]:
    return list(zip(periodos, periodos[1:]))


def _construir_aba_resumo(
    ws: Worksheet, empresa: str, periodos_selecionados: list[str], resumo_executivo: str
) -> None:
    ws.append([f"Empresa: {empresa}"])
    ws.append([f"Períodos analisados: {', '.join(periodos_selecionados)}"])
    ws.append([f"Data de geração: {datetime.now(timezone.utc).isoformat()}"])
    ws.append([])
    for linha in resumo_executivo.splitlines():
        ws.append([linha])


def _construir_aba_comparativo(
    ws: Worksheet, comparativo: list[ComparacaoItem], periodos_selecionados: list[str]
) -> None:
    pares = _pares_consecutivos(periodos_selecionados)
    cabecalho = ["Indicador", "Categoria", "Segmento", "Tipo de Período", "Ajustado/Reportado", "Unidade"]
    cabecalho += list(periodos_selecionados)
    cabecalho += [f"Variação Absoluta ({a}→{b})" for a, b in pares]
    cabecalho += [f"Variação % ou p.p. ({a}→{b})" for a, b in pares]
    cabecalho.append("Observações")
    ws.append(cabecalho)

    for item in comparativo:
        linha = [
            item.indicador,
            _valor_enum(item.categoria),
            _valor_enum(item.segmento),
            _valor_enum(item.tipo_periodo),
            _valor_enum(item.ajustado_ou_reportado),
            item.unidade,
        ]
        linha += [item.valores_por_periodo.get(periodo) for periodo in periodos_selecionados]
        linha += [item.variacao_absoluta.get(f"{a}_para_{b}") for a, b in pares]
        linha += [item.variacao_percentual_ou_pp.get(f"{a}_para_{b}") for a, b in pares]
        linha.append(item.observacoes)
        ws.append(linha)


def _construir_aba_evidencias(ws: Worksheet, indicadores: list[Indicador]) -> None:
    ws.append(
        [
            "Documento",
            "Página",
            "Indicador",
            "Valor Original",
            "Valor Normalizado",
            "Unidade",
            "Trecho-fonte",
            "Confiança",
            "Ajustado/Reportado",
            "Tipo de Período",
            "Segmento",
        ]
    )
    for indicador in indicadores:
        ws.append(
            [
                indicador.documento_id,
                indicador.pagina,
                indicador.indicador,
                indicador.valor_original,
                indicador.valor_normalizado,
                indicador.unidade,
                indicador.trecho_fonte,
                _valor_enum(indicador.confianca),
                _valor_enum(indicador.ajustado_ou_reportado),
                _valor_enum(indicador.tipo_periodo),
                _valor_enum(indicador.segmento),
            ]
        )


def _construir_aba_documentos(ws: Worksheet, documentos: list[Documento]) -> None:
    ws.append(
        ["Nome do documento", "Ano", "Trimestre", "URL de origem", "Data de download", "Número de páginas", "SHA256"]
    )
    for documento in documentos:
        ws.append(
            [
                documento.nome_arquivo,
                documento.ano,
                documento.trimestre,
                documento.url_origem,
                str(documento.data_download) if documento.data_download else None,
                documento.num_paginas,
                documento.sha256,
            ]
        )


def _construir_aba_pendencias(ws: Worksheet, pendencias: list[Pendencia]) -> None:
    ws.append(
        ["Descrição", "Tipo de gatilho", "Documento", "Página", "Indicador relacionado", "Trecho-fonte", "Ação sugerida"]
    )
    for pendencia in pendencias:
        ws.append(
            [
                pendencia.descricao,
                _valor_enum(pendencia.tipo_gatilho),
                pendencia.documento_id,
                pendencia.pagina,
                pendencia.indicador_relacionado,
                pendencia.trecho_fonte,
                pendencia.acao_sugerida,
            ]
        )


def _construir_aba_auditoria(ws: Worksheet, auditoria: list[ResultadoAuditoria]) -> None:
    ws.append(["Item verificado", "Resultado", "Detalhe", "Regra aplicada", "Data/hora"])
    for resultado in auditoria:
        ws.append(
            [
                resultado.item_verificado,
                _valor_enum(resultado.resultado),
                resultado.detalhe,
                resultado.regra_aplicada,
                resultado.timestamp.isoformat(),
            ]
        )


def construir_workbook(
    empresa: str,
    periodos_selecionados: list[str],
    documentos: list[Documento],
    indicadores: list[Indicador],
    comparativo: list[ComparacaoItem],
    resumo_executivo: str,
    pendencias: list[Pendencia],
    auditoria: list[ResultadoAuditoria],
) -> Workbook:
    wb = Workbook()

    resumo_ws = wb.active
    resumo_ws.title = "Resumo"
    _construir_aba_resumo(resumo_ws, empresa, periodos_selecionados, resumo_executivo)

    _construir_aba_comparativo(wb.create_sheet("Comparativo"), comparativo, periodos_selecionados)
    _construir_aba_evidencias(wb.create_sheet("Evidências"), indicadores)
    _construir_aba_documentos(wb.create_sheet("Documentos"), documentos)
    _construir_aba_pendencias(wb.create_sheet("Pendências"), pendencias)
    _construir_aba_auditoria(wb.create_sheet("Auditoria"), auditoria)

    return wb
