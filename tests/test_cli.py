import json
import re
from pathlib import Path

import pytest
import requests_mock

from magalu_analysis import config
from magalu_analysis.cli import ResumoAusenteError, executar_build_excel, executar_compare, executar_fetch
from magalu_analysis.summary_schema import ResumoInvalidoError

FIXTURES = Path(__file__).parent / "fixtures"
HTML_CENTRAL = (FIXTURES / "central_resultados.html").read_text(encoding="utf-8")
PDF_TEXTUAL = (FIXTURES / "sample_release.pdf").read_bytes()

INDICADOR_BASE = {
    "indicador": "Receita Líquida",
    "categoria": "financeiro",
    "segmento": "consolidado",
    "tipo_periodo": "trimestre",
    "ajustado_ou_reportado": None,
    "unidade": "R$ milhões",
    "pagina": 1,
    "trecho_fonte": "trecho",
    "confianca": "alta",
}


def _criar_run_com_indicadores(tmp_path, valores_por_documento):
    run_id = "run-teste"
    diretorio_run = tmp_path / run_id
    diretorio_documentos = diretorio_run / "documents"
    diretorio_documentos.mkdir(parents=True)

    documentos = []
    for doc_id, (ano, trimestre, valor) in valores_por_documento.items():
        documentos.append(
            {
                "documento_id": doc_id,
                "nome_arquivo": f"{doc_id}.pdf",
                "url_origem": "https://x/1.pdf",
                "ano": ano,
                "trimestre": trimestre,
                "data_download": "2026-09-17",
                "num_paginas": 2,
                "sha256": "a" * 64,
            }
        )
        (diretorio_documentos / f"{doc_id}.indicators.json").write_text(
            json.dumps([{**INDICADOR_BASE, "documento_id": doc_id, "valor_original": str(valor), "valor_normalizado": valor}]),
            encoding="utf-8",
        )

    manifest = {
        "run_id": run_id,
        "empresa": "Magazine Luiza",
        "n_releases": len(documentos),
        "periodos_selecionados": [[d["ano"], d["trimestre"]] for d in documentos],
        "data_execucao": "2026-09-17T10:00:00Z",
        "documentos": documentos,
    }
    (diretorio_run / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return run_id


def test_executar_fetch_seleciona_baixa_e_extrai_os_n_mais_recentes(tmp_path):
    with requests_mock.Mocker() as mock:
        mock.get(config.URL_CENTRAL_RESULTADOS, text=HTML_CENTRAL)
        mock.get(re.compile(r"^https://ri\.magazineluiza\.com\.br/Download\.aspx"), content=PDF_TEXTUAL)

        manifest = executar_fetch(n=2, diretorio_runs=tmp_path)

    assert manifest.n_releases == 2
    assert len(manifest.documentos) == 2
    assert manifest.periodos_selecionados == [(2026, 2), (2026, 1)]

    diretorio_run = tmp_path / manifest.run_id
    assert (diretorio_run / "manifest.json").exists()

    for documento in manifest.documentos:
        diretorio_documentos = diretorio_run / "documents"
        assert (diretorio_documentos / documento.nome_arquivo).exists()
        pages_path = diretorio_documentos / f"{documento.documento_id}.pages.json"
        assert pages_path.exists()
        paginas = json.loads(pages_path.read_text(encoding="utf-8"))
        assert len(paginas) == 2


def test_executar_fetch_registra_menos_documentos_quando_n_maior_que_disponivel(tmp_path):
    with requests_mock.Mocker() as mock:
        mock.get(config.URL_CENTRAL_RESULTADOS, text=HTML_CENTRAL)
        mock.get(re.compile(r"^https://ri\.magazineluiza\.com\.br/Download\.aspx"), content=PDF_TEXTUAL)

        manifest = executar_fetch(n=999, diretorio_runs=tmp_path)

    assert manifest.n_releases == 999
    assert len(manifest.documentos) < 999
    assert len(manifest.documentos) > 0


def test_executar_compare_gera_comparativo_json_e_retorna_series(tmp_path):
    run_id = _criar_run_com_indicadores(
        tmp_path, {"doc-q1": (2026, 1, 9205.7), "doc-q2": (2026, 2, 8898.7)}
    )

    comparativo = executar_compare(run_id, diretorio_runs=tmp_path)

    item = next(i for i in comparativo if i.indicador == "Receita Líquida")
    assert item.valores_por_periodo == {"1T26": 9205.7, "2T26": 8898.7}

    caminho_comparativo = tmp_path / run_id / "comparativo.json"
    assert caminho_comparativo.exists()
    dados = json.loads(caminho_comparativo.read_text(encoding="utf-8"))
    assert dados[0]["indicador"] == "Receita Líquida"


def test_executar_build_excel_falha_quando_resumo_executivo_esta_ausente(tmp_path):
    run_id = _criar_run_com_indicadores(tmp_path, {"doc-q1": (2026, 1, 9205.7)})
    executar_compare(run_id, diretorio_runs=tmp_path)

    with pytest.raises(ResumoAusenteError):
        executar_build_excel(run_id, diretorio_runs=tmp_path, diretorio_output=tmp_path / "output")


def test_executar_build_excel_falha_quando_resumo_tem_linguagem_proibida(tmp_path):
    run_id = _criar_run_com_indicadores(tmp_path, {"doc-q1": (2026, 1, 9205.7)})
    executar_compare(run_id, diretorio_runs=tmp_path)
    (tmp_path / run_id / "resumo_executivo.md").write_text(
        "Recomendamos a compra das ações.", encoding="utf-8"
    )

    with pytest.raises(ResumoInvalidoError):
        executar_build_excel(run_id, diretorio_runs=tmp_path, diretorio_output=tmp_path / "output")


def test_executar_build_excel_gera_arquivo_com_as_seis_abas(tmp_path):
    run_id = _criar_run_com_indicadores(
        tmp_path, {"doc-q1": (2026, 1, 9205.7), "doc-q2": (2026, 2, 8898.7)}
    )
    executar_compare(run_id, diretorio_runs=tmp_path)
    (tmp_path / run_id / "resumo_executivo.md").write_text(
        "A Receita Líquida caiu no trimestre.", encoding="utf-8"
    )

    from openpyxl import load_workbook

    destino = executar_build_excel(run_id, diretorio_runs=tmp_path, diretorio_output=tmp_path / "output")

    assert destino.exists()
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Resumo", "Comparativo", "Evidências", "Documentos", "Pendências", "Auditoria"]
    assert (tmp_path / run_id / "auditoria.json").exists()
